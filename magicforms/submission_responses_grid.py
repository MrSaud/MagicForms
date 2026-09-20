"""Build tabular submission answers for studio grid view and XLSX/PDF export."""

from __future__ import annotations

import os
import re
from datetime import date, timedelta
from io import BytesIO
from typing import Any, Iterable

from django.db.models import Count, Q
from django.db.models.functions import TruncDate
from django.utils import timezone
from django.utils.translation import gettext as _

from .i18n_db import gettext_db
from .timeline_i18n import timeline_event_kind_label, timeline_event_message_display, timeline_event_step_label
from .models import (
    FieldType,
    Form,
    FormField,
    FormSubmission,
    SubmissionEvent,
    SubmissionValue,
)
from .pdf_branding import stamp_entity_logo_on_pdf
from .pdf_text import (
    pdf_cell_paragraph,
    pdf_header_cell_paragraph,
    pdf_meta_paragraph,
    pdf_title_paragraph,
)

# Safety cap for export (all matching submissions).
MAX_EXPORT_ROWS = 5000
_CHECKLIST_VALUE_SCAN = 4000
_NUMBER_VALUE_SCAN = 8000


def display_submission_value(field: FormField, sv: SubmissionValue | None) -> str:
    """Plain-text cell for a field (aligned with print merge display)."""
    if sv is None:
        return ""
    raw = sv.value or ""
    ft = field.field_type
    if ft == FieldType.CHECKBOX:
        return "Yes" if raw == "yes" else "No"
    if ft == FieldType.CHECKLIST:
        lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
        return "; ".join(lines)
    if ft == FieldType.FILE:
        if sv.attachment:
            return os.path.basename(sv.attachment.name)
        return raw or ""
    return raw


def _truncate_cell(text: str, max_len: int) -> str:
    s = (text or "").replace("\r\n", "\n").strip()
    if len(s) <= max_len:
        return s
    return s[: max_len - 1] + "…"


def _value_count_rows(field: FormField, limit: int) -> list[dict[str, Any]]:
    return list(
        SubmissionValue.objects.filter(field=field)
        .exclude(value="")
        .values("value")
        .annotate(n=Count("id"))
        .order_by("-n")[:limit]
    )


def _field_chart_bar_h(field: FormField, rows: list[dict[str, Any]]) -> dict[str, Any] | None:
    if not rows:
        return None
    return {
        "canvas_id": f"mf-field-chart-{field.pk}",
        "title": field.label,
        "kind": "bar_h",
        "labels": [(r["value"] or "")[:200] for r in rows],
        "counts": [int(r["n"]) for r in rows],
    }


def _non_empty_answer_count(field: FormField, sqs) -> int:
    """Rows that count as ‘answered’ for completion charts (file = has attachment)."""
    if field.field_type == FieldType.FILE:
        return int(sqs.exclude(Q(attachment="") | Q(attachment__isnull=True)).count())
    return int(sqs.exclude(value="").count())


def _completion_doughnut(field: FormField, total_submissions: int) -> dict[str, Any]:
    sqs = SubmissionValue.objects.filter(field=field)
    n_ok = _non_empty_answer_count(field, sqs)
    n_miss = max(0, total_submissions - n_ok)
    return {
        "canvas_id": f"mf-field-chart-{field.pk}",
        "title": field.label,
        "kind": "doughnut",
        "labels": [str(_("With answer")), str(_("No answer"))],
        "counts": [n_ok, n_miss],
    }


def build_field_value_chart(field: FormField, total_submissions: int) -> dict[str, Any]:
    """One Chart.js payload per form field (distribution when possible, else completion)."""
    sqs = SubmissionValue.objects.filter(field=field)
    ft = field.field_type

    if ft in (FieldType.RADIO, FieldType.SELECT):
        rows = _value_count_rows(field, 12)
        bar = _field_chart_bar_h(field, rows)
        if bar is not None:
            return bar
        return _completion_doughnut(field, total_submissions)

    if ft == FieldType.CHECKBOX:
        n_tot = sqs.count()
        if n_tot == 0:
            return _completion_doughnut(field, total_submissions)
        n_yes = sqs.filter(value="yes").count()
        n_no = n_tot - n_yes
        return {
            "canvas_id": f"mf-field-chart-{field.pk}",
            "title": field.label,
            "kind": "doughnut",
            "labels": [str(_("Yes")), str(_("No"))],
            "counts": [n_yes, n_no],
        }

    if ft == FieldType.CHECKLIST:
        options = field.choice_list()
        if not options:
            return _completion_doughnut(field, total_submissions)
        counts: dict[str, int] = {o: 0 for o in options}
        for raw in sqs.exclude(value="").values_list("value", flat=True)[:_CHECKLIST_VALUE_SCAN]:
            seen: set[str] = set()
            for ln in str(raw or "").splitlines():
                s = ln.strip()
                if s in counts and s not in seen:
                    counts[s] += 1
                    seen.add(s)
        if not any(counts.values()):
            return _completion_doughnut(field, total_submissions)
        labels = list(options)
        data = [counts[o] for o in labels]
        return {
            "canvas_id": f"mf-field-chart-{field.pk}",
            "title": field.label,
            "kind": "bar_h",
            "labels": labels[:50],
            "counts": data[:50],
        }

    if ft == FieldType.NUMBER:
        vals: list[float] = []
        for raw in sqs.exclude(value="").values_list("value", flat=True)[:_NUMBER_VALUE_SCAN]:
            try:
                vals.append(float(str(raw).replace(",", ".").strip()))
            except (ValueError, TypeError):
                continue
        if not vals:
            return _completion_doughnut(field, total_submissions)
        mn, mx = min(vals), max(vals)
        n_bins = min(12, max(4, int(len(vals) ** 0.5)))
        if mn == mx:
            return {
                "canvas_id": f"mf-field-chart-{field.pk}",
                "title": field.label,
                "kind": "bar_v",
                "labels": [str(mn)],
                "counts": [len(vals)],
            }
        width = (mx - mn) / n_bins or 1.0
        bin_counts = [0] * n_bins
        for x in vals:
            i = min(int((x - mn) / width), n_bins - 1)
            bin_counts[i] += 1
        labels = []
        for i in range(n_bins):
            lo = mn + i * width
            hi = mn + (i + 1) * width
            labels.append(f"{lo:.4g}–{hi:.4g}")
        return {
            "canvas_id": f"mf-field-chart-{field.pk}",
            "title": field.label,
            "kind": "bar_v",
            "labels": labels,
            "counts": bin_counts,
        }

    if ft == FieldType.DATE:
        rows = _value_count_rows(field, 24)
        bar = _field_chart_bar_h(field, rows)
        if bar is not None:
            return bar
        return _completion_doughnut(field, total_submissions)

    if ft == FieldType.FILE:
        n_tot = sqs.count()
        if n_tot == 0:
            return _completion_doughnut(field, total_submissions)
        n_file = sqs.exclude(Q(attachment="") | Q(attachment__isnull=True)).count()
        n_empty = n_tot - n_file
        return {
            "canvas_id": f"mf-field-chart-{field.pk}",
            "title": field.label,
            "kind": "doughnut",
            "labels": [str(_("File uploaded")), str(_("No file"))],
            "counts": [n_file, n_empty],
        }

    if ft in (FieldType.TEXT, FieldType.TEXTAREA, FieldType.EMAIL):
        rows = _value_count_rows(field, 15)
        bar = _field_chart_bar_h(field, rows)
        if bar is not None:
            return bar
        return _completion_doughnut(field, total_submissions)

    return _completion_doughnut(field, total_submissions)


def build_responses_grid_statistics(form: Form, fields: list[FormField]) -> dict[str, Any]:
    """
    Aggregate counts for the responses grid overview: workflow, steps, last-30-day volume,
    daily spark data, and per-field answer distributions for Chart.js on the grid page.
    """
    subs = FormSubmission.objects.filter(form=form)
    total = subs.count()
    out: dict[str, Any] = {
        "total": total,
        "last_30d": 0,
        "unique_accounts": 0,
        "fields_count": len(fields),
        "workflow_bars": [],
        "step_bars": [],
        "daily_bars": [],
        "field_value_charts": [],
    }
    if total == 0:
        return out

    cutoff = timezone.now() - timedelta(days=30)
    out["last_30d"] = subs.filter(submitted_at__gte=cutoff).count()
    out["unique_accounts"] = subs.filter(submitted_by__isnull=False).values("submitted_by_id").distinct().count()

    for state, label in FormSubmission.WorkflowState.choices:
        n = subs.filter(workflow_state=state).count()
        out["workflow_bars"].append(
            {
                "key": state,
                "label": str(label),
                "count": n,
                "pct": round(100.0 * n / total, 1),
            }
        )

    step_rows = list(
        subs.values("current_step__label").annotate(n=Count("id")).order_by("-n")[:16]
    )
    max_step = max((r["n"] for r in step_rows), default=1)
    for r in step_rows:
        raw_lbl = r["current_step__label"]
        lbl = str(_("No step")) if not (raw_lbl or "").strip() else raw_lbl
        out["step_bars"].append(
            {
                "label": lbl,
                "count": r["n"],
                "pct_width": round(100.0 * r["n"] / max_step, 1),
            }
        )

    tz = timezone.get_current_timezone()
    end_d = timezone.localdate()
    start_d = end_d - timedelta(days=29)
    daily_rows = subs.annotate(day=TruncDate("submitted_at", tzinfo=tz)).values("day").annotate(n=Count("id"))
    day_map: dict[date, int] = {}
    for row in daily_rows:
        d = row["day"]
        if d is None:
            continue
        if isinstance(d, str):
            try:
                d = date.fromisoformat(d[:10])
            except ValueError:
                continue
        elif not isinstance(d, date):
            continue
        day_map[d] = row["n"]
    max_day = max(day_map.values(), default=1)
    for i in range(30):
        d = start_d + timedelta(days=i)
        n = day_map.get(d, 0)
        out["daily_bars"].append(
            {
                "date": d.isoformat(),
                "label": f"{d.day} {d.strftime('%b')}",
                "count": n,
                "pct_height": round(100.0 * n / max_day, 1) if max_day else 0.0,
            }
        )

    out["field_value_charts"] = [build_field_value_chart(ff, total) for ff in fields]

    return out


def build_grid_rows(
    submissions: Iterable[FormSubmission],
    fields: list[FormField],
) -> tuple[list[list[str]], list[str]]:
    """
    Return ``(data_rows, field_keys)`` where each inner list is fixed columns then field values
    in ``fields`` order. ``field_keys`` are ``str(field.pk)`` for template alignment.
    """
    field_keys = [str(f.pk) for f in fields]
    rows: list[list[str]] = []
    for sub in submissions:
        by_field: dict[int, SubmissionValue] = {v.field_id: v for v in sub.values.all()}
        ref = str(sub.reference_token)
        ts = timezone.localtime(sub.submitted_at).strftime("%Y-%m-%d %H:%M")
        email = (sub.submitter_email or "").strip()
        account = sub.submitted_by.get_username() if sub.submitted_by_id else ""
        wf = sub.get_workflow_state_display()
        step = sub.current_step.label if sub.current_step_id and sub.current_step else ""
        cells = [ref, ts, email, account, wf, step]
        for f in fields:
            cells.append(display_submission_value(f, by_field.get(f.pk)))
        rows.append(cells)
    return rows, field_keys


def export_xlsx_bytes(headers: list[str], data_rows: list[list[str]], sheet_title: str) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is not installed.") from exc

    wb = Workbook()
    ws = wb.active
    safe_title = re.sub(r"[\[\]\:\*\?\/\\]", "_", (sheet_title or "Responses")[:31]) or "Responses"
    ws.title = safe_title

    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    for r, row in enumerate(data_rows, start=2):
        for cc, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=cc, value=_truncate_cell(str(val), 32700))
            cell.number_format = "@"  # force text to avoid Excel mangling UUIDs / dates

    for c in range(1, len(headers) + 1):
        col_letter = get_column_letter(c)
        ws.column_dimensions[col_letter].width = min(48, max(10, len(headers[c - 1]) + 2))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf_bytes(
    title: str,
    headers: list[str],
    data_rows: list[list[str]],
    *,
    entity=None,
) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, letter
        from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise RuntimeError("reportlab is not installed.") from exc

    buf = BytesIO()
    page = landscape(letter)
    doc = SimpleDocTemplate(
        buf,
        pagesize=page,
        leftMargin=28,
        rightMargin=28,
        topMargin=36,
        bottomMargin=36,
    )
    story = [pdf_title_paragraph(title), Spacer(1, 10)]

    safe_headers = [_truncate_cell(h, 80) for h in headers]
    safe_rows = [[_truncate_cell(str(v), 400) for v in row] for row in data_rows]
    table_data = [
        [pdf_header_cell_paragraph(h, font_size=7) for h in safe_headers],
        *[
            [pdf_cell_paragraph(v, font_size=6, leading=8) for v in row]
            for row in safe_rows
        ],
    ]

    ncols = len(safe_headers)
    avail = page[0] - 56
    col_w = max(36, min(120, avail / max(1, ncols)))

    tbl = Table(table_data, colWidths=[col_w] * ncols, repeatRows=1)
    tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0071e3")),
                ("FONTSIZE", (0, 0), (-1, 0), 7),
                ("FONTSIZE", (0, 1), (-1, -1), 6),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.Color(0.97, 0.97, 0.98), colors.white]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c8c8cc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story.append(tbl)
    doc.build(story)
    raw = buf.getvalue()
    if entity is not None:
        raw = stamp_entity_logo_on_pdf(raw, entity)
    return raw


def _timeline_actor_label(user) -> str:
    if not user:
        return ""
    by = user.get_username()
    fn = (user.get_full_name() or "").strip()
    if fn:
        return f"{fn} ({by})"
    return by


def _timeline_pdf_date_display(at) -> str:
    """Match studio ``mf-timeline-date`` (``M j, Y, P`` + uppercase)."""
    from django.utils.formats import date_format

    s = date_format(
        timezone.localtime(at),
        format="M j, Y, P",
        use_l10n=True,
    )
    return s.upper()


def _timeline_pdf_actor_display(user) -> str:
    """Match manage timeline template ``{{ ev.created_by }}`` (username)."""
    if not user:
        return ""
    return (user.get_username() or "").strip()


def build_submission_timeline_pdf_payload(
    submission: FormSubmission,
) -> tuple[str, list[str], list[dict[str, str]]]:
    """
    Title, optional summary lines, and timeline entries for :func:`export_timeline_pdf_bytes`.

    Same data as the studio ``mf-timeline`` list (submission events only, chronological).
    """
    from django.utils.formats import date_format

    form = submission.form
    ent = form.entity
    entity_title = gettext_db(str(ent.get_public_display_title()))
    form_title = gettext_db(str(form.title))
    title = f"{entity_title} — {gettext_db(_('Submission timeline'))}"

    submitted_local = timezone.localtime(submission.submitted_at)
    submitted_str = date_format(
        submitted_local,
        format="SHORT_DATETIME_FORMAT",
        use_l10n=True,
    )
    wf = gettext_db(str(submission.get_workflow_state_display()))
    step_now = ""
    if submission.current_step_id and submission.current_step:
        step_now = gettext_db(str(submission.current_step.label))

    applicant = ""
    if submission.submitted_by_id:
        applicant = _timeline_actor_label(submission.submitted_by)
    elif (submission.submitter_email or "").strip():
        applicant = (submission.submitter_email or "").strip()

    summary = [
        f"{gettext_db(_('Reference'))}: {submission.reference_token}",
        f"{gettext_db(_('Form'))}: {form_title}",
        f"{gettext_db(_('Workflow'))}: {wf}"
        + (f" · {gettext_db(_('Current step'))}: {step_now}" if step_now else ""),
        f"{gettext_db(_('Submitted'))}: {submitted_str}",
    ]
    if applicant:
        summary.append(f"{gettext_db(_('Applicant'))}: {applicant}")

    entries: list[dict[str, str]] = []
    for ev in (
        SubmissionEvent.objects.filter(submission_id=submission.pk)
        .select_related("step", "created_by")
        .order_by("created_at")
    ):
        step_label = str(ev.step.label) if ev.step_id else None
        entries.append(
            {
                "date": _timeline_pdf_date_display(ev.created_at),
                "title": timeline_event_kind_label(ev),
                "step": timeline_event_step_label(ev.step) if ev.step_id else "",
                "actor": _timeline_pdf_actor_display(ev.created_by),
                "message": timeline_event_message_display(
                    ev.message,
                    step_label=step_label,
                ),
            }
        )

    if not entries:
        entries.append(
            {
                "date": "",
                "title": gettext_db(_("No events.")),
                "step": "",
                "actor": "",
                "message": "",
            }
        )

    return title, summary, entries


def _timeline_pdf_item_flowables(entry: dict[str, str], *, content_width: float) -> list:
    """One ``mf-timeline-item`` — same lines as studio ``mf-timeline`` (full width, no clipping)."""
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Spacer

    from .pdf_text import pdf_font_names, pdf_paragraph_markup

    reg_font, bold_font = pdf_font_names()
    muted = colors.HexColor("#5a5f6a")
    body_color = colors.HexColor("#1d1d1f")
    left = 12
    uid = id(entry)
    rail = colors.Color(0, 113 / 255, 227 / 255, alpha=0.25)

    def ps(
        suffix: str,
        font: str,
        size: float,
        leading: float,
        color,
        *,
        rail_border: bool = False,
    ) -> ParagraphStyle:
        kw: dict = dict(
            name=f"TimelinePdf_{uid}_{suffix}",
            fontName=font,
            fontSize=size,
            leading=leading,
            textColor=color,
            leftIndent=left,
            rightIndent=0,
            spaceBefore=0,
            spaceAfter=2,
            wordWrap="CJK",
        )
        if rail_border:
            kw["borderWidth"] = (2, 0, 0, 0)
            kw["borderColor"] = rail
            kw["borderPadding"] = (8, 0, 4, 0)
        return ParagraphStyle(**kw)

    lines: list = []
    first_line = True
    if entry.get("date"):
        lines.append(
            Paragraph(
                pdf_paragraph_markup(entry["date"]),
                ps("date", reg_font, 7.5, 9, muted, rail_border=first_line),
            )
        )
        first_line = False
    title = (entry.get("title") or "").strip()
    if title:
        lines.append(
            Paragraph(
                pdf_paragraph_markup(title),
                ps("title", bold_font, 10, 12, body_color, rail_border=first_line),
            )
        )
        first_line = False
    step = (entry.get("step") or "").strip()
    if step:
        lines.append(
            Paragraph(
                pdf_paragraph_markup(step),
                ps("step", reg_font, 9, 11, muted),
            )
        )
    actor = (entry.get("actor") or "").strip()
    if actor:
        lines.append(
            Paragraph(
                pdf_paragraph_markup(f"· {actor}"),
                ps("actor", reg_font, 9, 11, muted),
            )
        )
    message = (entry.get("message") or "").strip()
    if message:
        lines.append(
            Paragraph(
                pdf_paragraph_markup(message),
                ps("msg", reg_font, 9, 12, body_color),
            )
        )

    if not lines:
        return [Spacer(1, 14)]

    lines.append(Spacer(1, 14))
    return lines


def export_timeline_pdf_bytes(
    document_title: str,
    entries: list[dict[str, str]],
    *,
    summary_lines: list[str] | None = None,
    entity=None,
) -> bytes:
    """
    Portrait PDF matching the studio timeline list (not a table).
    """
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise RuntimeError("reportlab is not installed.") from exc

    from .pdf_text import pdf_font_names, pdf_paragraph_markup

    has_logo = bool(
        entity is not None
        and getattr(entity, "page_logo", None)
        and entity.page_logo
    )
    buf = BytesIO()
    page = letter
    left_margin = 54
    right_margin = 54
    content_width = page[0] - left_margin - right_margin
    doc = SimpleDocTemplate(
        buf,
        pagesize=page,
        leftMargin=left_margin,
        rightMargin=right_margin,
        topMargin=84 if has_logo else 54,
        bottomMargin=54,
    )
    story: list = []
    bold_font = pdf_font_names()[1]
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph

    timeline_heading = ParagraphStyle(
        name="TimelinePdfHeading",
        fontName=bold_font,
        fontSize=14,
        leading=17,
        spaceAfter=12,
        wordWrap="CJK",
    )
    story.append(
        Paragraph(pdf_paragraph_markup(gettext_db(_("Timeline"))), timeline_heading)
    )

    for entry in entries:
        story.extend(_timeline_pdf_item_flowables(entry, content_width=content_width))

    doc.build(story)
    raw = buf.getvalue()
    if entity is not None:
        raw = stamp_entity_logo_on_pdf(raw, entity)
    return raw
